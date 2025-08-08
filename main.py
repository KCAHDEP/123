# -*- coding: utf-8 -*-
"""
Notification Generator - GUI for Windows (PyQt5)
Features:
- Paste template text (with placeholders {{flat}}, {{date}}, {{time}})
- Paste list of apartment numbers (one per line or separated by spaces/commas)
- Choose date range and time range, generate random datetime per apartment
- Saves settings/history to JSON, outputs .docx files and a ZIP archive
"""
import sys
import os
import json
import random
from datetime import datetime, timedelta, time as dtime
from pathlib import Path

from PyQt5.QtWidgets import (QApplication, QWidget, QVBoxLayout, QHBoxLayout,
                             QLabel, QPushButton, QTextEdit, QLineEdit,
                             QFileDialog, QDateEdit, QTimeEdit, QMessageBox)
from PyQt5.QtCore import Qt, QDate, QTime

from docx import Document
from zipfile import ZipFile

APP_DIR = Path.home() / "NotificationGenerator"
APP_DIR.mkdir(exist_ok=True)

SETTINGS_FILE = APP_DIR / "settings.json"
HISTORY_FILE = APP_DIR / "history.json"

def load_settings():
    if SETTINGS_FILE.exists():
        try:
            with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            return {}
    return {}

def save_settings(data):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def append_history(entry):
    history = []
    if HISTORY_FILE.exists():
        try:
            with open(HISTORY_FILE, "r", encoding='utf-8') as f:
                history = json.load(f)
        except:
            history = []
    history.append(entry)
    with open(HISTORY_FILE, "w", encoding='utf-8') as f:
        json.dump(history, f, ensure_ascii=False, indent=2)

def parse_apartments(text):
    parts = []
    for token in text.replace(",", " ").split():
        token = token.strip()
        if token == "":
            continue
        num = "".join(ch for ch in token if ch.isdigit())
        if num:
            parts.append(num)
    seen = set()
    uniq = []
    for p in parts:
        if p not in seen:
            seen.add(p)
            uniq.append(p)
    return uniq

def random_datetime_between(start_date, end_date, start_time, end_time):
    days_delta = (end_date - start_date).days
    rand_day = random.randint(0, max(0, days_delta))
    chosen_date = start_date + timedelta(days=rand_day)
    st_seconds = start_time.hour * 3600 + start_time.minute * 60 + start_time.second
    en_seconds = end_time.hour * 3600 + end_time.minute * 60 + end_time.second
    if en_seconds < st_seconds:
        en_seconds = st_seconds
    rand_seconds = random.randint(st_seconds, en_seconds)
    hours = rand_seconds // 3600
    mins = (rand_seconds % 3600) // 60
    return datetime(chosen_date.year, chosen_date.month, chosen_date.day, hours, mins)

def make_docx_from_text(text, out_path):
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(str(out_path))

from PyQt5.QtWidgets import QTimeEdit

class MainWindow(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Генератор уведомлений")
        self.resize(900, 700)
        self.settings = load_settings()
        self._build_ui()
        self._load_settings_into_ui()

    def _build_ui(self):
        layout = QVBoxLayout()
        layout.addWidget(QLabel("Шаблон уведомления (используйте {{flat}}, {{date}}, {{time}}):"))
        self.template_edit = QTextEdit()
        layout.addWidget(self.template_edit)
        btns_h = QHBoxLayout()
        self.load_template_btn = QPushButton("Загрузить шаблон из файла")
        self.save_template_btn = QPushButton("Сохранить шаблон в файл")
        btns_h.addWidget(self.load_template_btn)
        btns_h.addWidget(self.save_template_btn)
        layout.addLayout(btns_h)
        self.load_template_btn.clicked.connect(self.load_template_from_file)
        self.save_template_btn.clicked.connect(self.save_template_to_file)

        layout.addWidget(QLabel("Список номеров квартир (через пробел/запятую или построчно):"))
        self.apts_edit = QTextEdit()
        layout.addWidget(self.apts_edit)
        apts_btn_h = QHBoxLayout()
        self.load_apts_btn = QPushButton("Загрузить список из файла")
        apts_btn_h.addWidget(self.load_apts_btn)
        layout.addLayout(apts_btn_h)
        self.load_apts_btn.clicked.connect(self.load_apartments_from_file)

        dr_layout = QHBoxLayout()
        dr_layout.addWidget(QLabel("Дата от:"))
        self.date_from = QDateEdit(calendarPopup=True)
        self.date_from.setDate(QDate.currentDate())
        dr_layout.addWidget(self.date_from)
        dr_layout.addWidget(QLabel("по:"))
        self.date_to = QDateEdit(calendarPopup=True)
        self.date_to.setDate(QDate.currentDate().addDays(3))
        dr_layout.addWidget(self.date_to)

        tr_layout = QHBoxLayout()
        tr_layout.addWidget(QLabel("Время от:"))
        self.time_from = QTimeEdit()
        self.time_from.setTime(QTime(8,0))
        tr_layout.addWidget(self.time_from)
        tr_layout.addWidget(QLabel("по:"))
        self.time_to = QTimeEdit()
        self.time_to.setTime(QTime(17,0))
        tr_layout.addWidget(self.time_to)

        layout.addLayout(dr_layout)
        layout.addLayout(tr_layout)

        out_h = QHBoxLayout()
        out_h.addWidget(QLabel("Имя архива (без расширения):"))
        self.archive_name = QLineEdit("уведомления")
        out_h.addWidget(self.archive_name)
        layout.addLayout(out_h)

        gen_h = QHBoxLayout()
        self.generate_btn = QPushButton("Сгенерировать уведомления и создать ZIP")
        gen_h.addWidget(self.generate_btn)
        layout.addLayout(gen_h)
        self.generate_btn.clicked.connect(self.on_generate)

        self.status_label = QLabel("")
        layout.addWidget(self.status_label)

        settings_h = QHBoxLayout()
        self.save_settings_btn = QPushButton("Сохранить настройки")
        self.load_settings_btn = QPushButton("Загрузить настройки")
        settings_h.addWidget(self.save_settings_btn)
        settings_h.addWidget(self.load_settings_btn)
        layout.addLayout(settings_h)
        self.save_settings_btn.clicked.connect(self.save_settings_action)
        self.load_settings_btn.clicked.connect(self.load_settings_action)

        self.setLayout(layout)

    def load_template_from_file(self):
        fname, _ = QFileDialog.getOpenFileName(self, "Открыть файл шаблона", "", "Текстовые файлы (*.txt);;Все файлы (*)")
        if fname:
            try:
                with open(fname, "r", encoding="utf-8") as f:
                    data = f.read()
                self.template_edit.setPlainText(data)
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Не удалось прочитать файл: {e}")

    def save_template_to_file(self):
        fname, _ = QFileDialog.getSaveFileName(self, "Сохранить шаблон как", "template.txt", "Текстовые файлы (*.txt);;Все файлы (*)")
        if fname:
            try:
                with open(fname, "w", encoding="utf-8") as f:
                    f.write(self.template_edit.toPlainText())
                QMessageBox.information(self, "Сохранено", "Шаблон сохранён.")
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить файл: {e}")

    def load_apartments_from_file(self):
        fname, _ = QFileDialog.getOpenFileName(self, "Открыть файл со списком квартир", "", "Текстовые файлы (*.txt);;Все файлы (*)")
        if fname:
            try:
                with open(fname, "r", encoding="utf-8") as f:
                    data = f.read()
                self.apts_edit.setPlainText(data)
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Не удалось прочитать файл: {e}")

    def save_settings_action(self):
        data = {
            "template": self.template_edit.toPlainText(),
            "apartments": self.apts_edit.toPlainText(),
            "date_from": self.date_from.date().toString(Qt.ISODate),
            "date_to": self.date_to.date().toString(Qt.ISODate),
            "time_from": self.time_from.time().toString(),
            "time_to": self.time_to.time().toString(),
            "archive_name": self.archive_name.text()
        }
        save_settings(data)
        QMessageBox.information(self, "Сохранено", f"Настройки сохранены в:\\n{SETTINGS_FILE}")

    def load_settings_action(self):
        s = load_settings()
        if not s:
            QMessageBox.information(self, "Настройки", "Файл настроек не найден.")
            return
        self.template_edit.setPlainText(s.get("template", ""))
        self.apts_edit.setPlainText(s.get("apartments", ""))
        try:
            if s.get("date_from"):
                qd = QDate.fromString(s.get("date_from"), Qt.ISODate)
                self.date_from.setDate(qd)
            if s.get("date_to"):
                qd = QDate.fromString(s.get("date_to"), Qt.ISODate)
                self.date_to.setDate(qd)
            if s.get("time_from"):
                qt = QTime.fromString(s.get("time_from"))
                if qt.isValid():
                    self.time_from.setTime(qt)
            if s.get("time_to"):
                qt = QTime.fromString(s.get("time_to"))
                if qt.isValid():
                    self.time_to.setTime(qt)
            self.archive_name.setText(s.get("archive_name", "уведомления"))
        except Exception:
            pass
        QMessageBox.information(self, "Настройки", "Настройки загружены.")

    def on_generate(self):
        template_text = self.template_edit.toPlainText().strip()
        if not template_text:
            QMessageBox.warning(self, "Шаблон пуст", "Пожалуйста, введите шаблон уведомления.")
            return
        apts_text = self.apts_edit.toPlainText()
        apartments = parse_apartments(apts_text)
        if not apartments:
            QMessageBox.warning(self, "Квартиры не заданы", "Пожалуйста, введите список квартир.")
            return
        dt_from = self.date_from.date().toPyDate()
        dt_to = self.date_to.date().toPyDate()
        t_from_q = self.time_from.time()
        t_to_q = self.time_to.time()
        t_from = dtime(t_from_q.hour(), t_from_q.minute())
        t_to = dtime(t_to_q.hour(), t_to_q.minute())
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_folder = Path.cwd() / f"output_notifications_{timestamp}"
        out_folder.mkdir(parents=True, exist_ok=True)
        generated_files = []
        for apt in apartments:
            dt = random_datetime_between(dt_from, dt_to, t_from, t_to)
            date_str = dt.strftime("%d.%m.%Y")
            time_str = dt.strftime("%H:%M")
            text = template_text
            text = text.replace("ЖК Салют", "ЖК Красный Металлист")
            text = text.replace("жк Салют", "ЖК Красный Металлист")
            text = text.replace("жк салют", "ЖК Красный Металлист")
            text = text.replace("ул. 50 лет ВЛКСМ, д. 11/1", "ул. Гражданская, д. 1/1")
            text = text.replace("Квартира № 19", f"Квартира № {apt}")
            text = text.replace("{{flat}}", str(apt))
            text = text.replace("{{date}}", date_str)
            text = text.replace("{{time}}", time_str)
            filename = f"Уведомление_кв_{apt}.docx"
            out_path = out_folder / filename
            try:
                make_docx_from_text(text, out_path)
                generated_files.append(out_path)
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Не удалось создать {filename}:\\n{e}")
                return
        archive_name = self.archive_name.text().strip() or f"уведомления_{timestamp}"
        zip_path = Path.cwd() / f"{archive_name}.zip"
        with ZipFile(zip_path, "w") as zf:
            for f in generated_files:
                zf.write(f, arcname=f.name)
        entry = {"timestamp": datetime.now().isoformat(), "count": len(generated_files), "archive": str(zip_path)}
        append_history(entry)
        QMessageBox.information(self, "Готово", f"Сгенерировано {len(generated_files)} уведомлений.\\nАрхив: {zip_path}")
        self.status_label.setText(f"Сгенерировано {len(generated_files)} файлов. Архив: {zip_path}")

def main():
    app = QApplication(sys.argv)
    w = MainWindow()
    w.show()
    sys.exit(app.exec_())

if __name__ == "__main__":
    main()
